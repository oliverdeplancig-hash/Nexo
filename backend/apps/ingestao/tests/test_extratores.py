import io

import pytest
from docx import Document

from apps.ingestao.extratores import FalhaExtracao, extrair
from apps.ingestao.fatiamento import fatiar, sugerir_palavras_chave, sugerir_referencia


def test_docx_preserva_secoes():
    doc = Document()
    doc.add_heading("Capítulo 1 — Prazos", level=1)
    doc.add_paragraph("O pedido deve ser feito em 30 dias.")
    doc.add_heading("Capítulo 2 — Documentos", level=1)
    doc.add_paragraph("Apresente o formulário assinado.")
    buffer = io.BytesIO()
    doc.save(buffer)
    blocos = extrair(buffer.getvalue(), ".docx")
    assert [b.secao for b in blocos] == ["Capítulo 1 — Prazos", "Capítulo 2 — Documentos"]


def test_html_ignora_scripts():
    blocos = extrair(b"<h1>Guia</h1><script>alert(1)</script><p>Texto do guia.</p>", ".html")
    assert "alert" not in " ".join(b.texto for b in blocos)
    assert blocos[0].secao == "Guia"


def test_texto_vazio_gera_falha_amigavel():
    with pytest.raises(FalhaExtracao, match="OCR"):
        extrair(b"   \n  ", ".txt")


def test_fatiamento_nao_atravessa_pagina_e_e_estavel():
    from apps.ingestao.extratores import Bloco

    blocos = [Bloco(1, "", "Parágrafo um. " * 200), Bloco(2, "", "Parágrafo dois. " * 10)]
    a = fatiar(blocos, 500, 100)
    assert a == fatiar(blocos, 500, 100)
    assert {t.pagina for t in a} == {1, 2}
    assert all(len(t.texto) <= 500 for t in a)


def test_sugestoes_de_metadados():
    texto = "Portaria nº 123/2024. Dispõe sobre férias. As férias dos servidores..."
    assert sugerir_referencia(texto).startswith("Portaria nº 123/2024")
    assert sugerir_palavras_chave(texto)[0] == "ferias"
