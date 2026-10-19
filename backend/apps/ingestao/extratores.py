"""Extração de texto por formato. Cada extrator devolve blocos com página e seção."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass
from html.parser import HTMLParser

from .ocr import OCRIndisponivel, imagens_de_arquivo, obter_ocr, renderizar_pagina_pdf


class FalhaExtracao(Exception):
    """Mensagem já vem pronta para o usuário."""


@dataclass
class Bloco:
    pagina: int | None
    secao: str
    texto: str
    ocr: bool = False


# Página com menos caracteres úteis que isto é tratada como digitalizada.
MINIMO_CARACTERES_PAGINA = 25
EXTENSOES_IMAGEM = (".png", ".jpg", ".jpeg", ".tif", ".tiff")


_TITULO_MD = re.compile(r"^(#{1,4})\s+(.+)$")
_TITULO_NORMATIVO = re.compile(
    r"^\s*((CAP[IÍ]TULO|T[IÍ]TULO|SE[CÇ][AÃ]O|ANEXO)\s+[IVXLC\d]+.*|Art\.?\s*\d+[ºo°]?.{0,80})$",
    flags=re.IGNORECASE,
)


def verificar_assinatura(conteudo: bytes, extensao: str) -> None:
    """Confere a assinatura do arquivo com a extensão declarada."""
    if extensao == ".pdf" and not conteudo.startswith(b"%PDF"):
        raise FalhaExtracao("O arquivo não é um PDF válido ou está corrompido.")
    if extensao == ".docx" and not conteudo.startswith(b"PK"):
        raise FalhaExtracao("O arquivo não é um DOCX válido ou está corrompido.")
    assinaturas_imagem = {
        ".png": (b"\x89PNG",),
        ".jpg": (b"\xff\xd8",),
        ".jpeg": (b"\xff\xd8",),
        ".tif": (b"II*\x00", b"MM\x00*"),
        ".tiff": (b"II*\x00", b"MM\x00*"),
    }
    if extensao in assinaturas_imagem and not conteudo.startswith(assinaturas_imagem[extensao]):
        raise FalhaExtracao("O arquivo não é uma imagem válida ou está corrompido.")


def extrair(conteudo: bytes, extensao: str) -> list[Bloco]:
    verificar_assinatura(conteudo, extensao)
    try:
        if extensao == ".pdf":
            blocos = _pdf(conteudo)
        elif extensao in EXTENSOES_IMAGEM:
            blocos = _imagem(conteudo)
        elif extensao == ".docx":
            blocos = _docx(conteudo)
        elif extensao in (".txt", ".md"):
            blocos = _texto(_decodificar(conteudo), markdown=extensao == ".md")
        elif extensao in (".html", ".htm"):
            blocos = _html(_decodificar(conteudo))
        else:
            raise FalhaExtracao(f"Formato não suportado: {extensao}")
    except FalhaExtracao:
        raise
    except OCRIndisponivel as erro:
        raise FalhaExtracao(f"Não foi possível aplicar OCR: {erro}") from erro
    except Exception as erro:  # bibliotecas de terceiros lançam exceções variadas
        raise FalhaExtracao(f"Não foi possível ler o arquivo ({erro.__class__.__name__}).") from erro

    blocos = [b for b in blocos if b.texto.strip()]
    if not blocos:
        if obter_ocr().disponivel():
            raise FalhaExtracao("Nenhum texto pôde ser extraído, nem por OCR. Verifique a qualidade da digitalização.")
        raise FalhaExtracao(
            "Nenhum texto pôde ser extraído. O documento parece digitalizado (imagem) e o OCR "
            "não está disponível neste computador. Instale o Tesseract com o idioma português "
            "(ver docs/architecture/implantacao.md) e use 'Tentar novamente'."
        )
    return blocos


def _decodificar(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise FalhaExtracao("Codificação de texto não reconhecida.")


def _textos_pdfium(conteudo: bytes) -> list[str] | None:
    """Texto de cada página pelo pdfium, que preserva os espaços entre palavras em PDFs onde o pypdf os perde
    ("POLÍCIAMILITARDABAHIA"), o que prejudica a busca (ADR-025). None se o pdfium não abrir o arquivo."""
    try:
        import pypdfium2 as pdfium

        documento = pdfium.PdfDocument(conteudo)
    except Exception:
        return None
    try:
        textos = []
        for indice in range(len(documento)):
            pagina = documento[indice]
            pagina_texto = pagina.get_textpage()
            try:
                bruto = pagina_texto.get_text_range()
            finally:
                pagina_texto.close()
                pagina.close()
            textos.append(bruto.replace("\r\n", "\n").replace("\r", "\n").replace("\ufffe", "-"))
        return textos
    except Exception:
        return None
    finally:
        documento.close()


def _textos_pypdf(conteudo: bytes) -> list[str]:
    from pypdf import PdfReader

    leitor = PdfReader(io.BytesIO(conteudo))
    if leitor.is_encrypted:
        try:
            leitor.decrypt("")
        except Exception as erro:
            raise FalhaExtracao("O PDF está protegido por senha.") from erro
    return [pagina.extract_text() or "" for pagina in leitor.pages]


def _pdf(conteudo: bytes) -> list[Bloco]:
    textos = _textos_pdfium(conteudo)
    if textos is None:
        textos = _textos_pypdf(conteudo)
    blocos: list[Bloco] = []
    secao = ""
    ocr = obter_ocr()
    for numero, texto in enumerate(textos, start=1):
        texto = texto or ""
        por_ocr = False
        if len("".join(texto.split())) < MINIMO_CARACTERES_PAGINA and ocr.disponivel():
            texto = ocr.reconhecer(renderizar_pagina_pdf(conteudo, numero - 1))
            por_ocr = True
        for linha in texto.splitlines():
            if _TITULO_NORMATIVO.match(linha.strip()) and len(linha.strip()) < 120:
                secao = linha.strip()
                break
        blocos.append(Bloco(pagina=numero, secao=secao, texto=texto, ocr=por_ocr))
    return blocos


def _imagem(conteudo: bytes) -> list[Bloco]:
    ocr = obter_ocr()
    if not ocr.disponivel():
        raise FalhaExtracao(
            "Imagens dependem de OCR, que não está disponível neste computador. Instale o Tesseract "
            "com o idioma português (ver docs/architecture/implantacao.md)."
        )
    return [
        Bloco(pagina=numero, secao="", texto=ocr.reconhecer(imagem), ocr=True)
        for numero, imagem in enumerate(imagens_de_arquivo(conteudo), start=1)
    ]


def _docx(conteudo: bytes) -> list[Bloco]:
    from docx import Document

    doc = Document(io.BytesIO(conteudo))
    blocos: list[Bloco] = []
    secao = ""
    atual: list[str] = []

    def _fechar():
        if atual:
            blocos.append(Bloco(pagina=None, secao=secao, texto="\n".join(atual)))
            atual.clear()

    for paragrafo in doc.paragraphs:
        estilo = (paragrafo.style.name or "").lower() if paragrafo.style is not None else ""
        texto = paragrafo.text.strip()
        if not texto:
            continue
        if estilo.startswith(("heading", "título", "titulo", "title")):
            _fechar()
            secao = texto[:255]
        atual.append(texto)
    _fechar()
    for tabela in doc.tables:
        linhas = [" | ".join(c.text.strip() for c in linha.cells) for linha in tabela.rows]
        if linhas:
            blocos.append(Bloco(pagina=None, secao=secao or "Tabela", texto="\n".join(linhas)))
    return blocos


def _texto(conteudo: str, markdown: bool) -> list[Bloco]:
    blocos: list[Bloco] = []
    secao = ""
    atual: list[str] = []
    for linha in conteudo.splitlines():
        m = _TITULO_MD.match(linha) if markdown else None
        if m:
            if atual:
                blocos.append(Bloco(None, secao, "\n".join(atual)))
                atual = []
            secao = m.group(2).strip()[:255]
        atual.append(linha)
    if atual:
        blocos.append(Bloco(None, secao, "\n".join(atual)))
    return blocos


class _ExtratorHTML(HTMLParser):
    IGNORAR = {"script", "style", "noscript", "head"}
    TITULOS = {"h1", "h2", "h3"}

    def __init__(self):
        super().__init__()
        self.blocos: list[Bloco] = []
        self.secao = ""
        self.atual: list[str] = []
        self._ignorando = 0
        self._em_titulo = False

    def handle_starttag(self, tag, attrs):
        if tag in self.IGNORAR:
            self._ignorando += 1
        elif tag in self.TITULOS:
            self._fechar()
            self._em_titulo = True
            self.secao = ""

    def handle_endtag(self, tag):
        if tag in self.IGNORAR and self._ignorando:
            self._ignorando -= 1
        elif tag in self.TITULOS:
            self._em_titulo = False
        elif tag in ("p", "li", "tr", "div", "br"):
            self.atual.append("\n")

    def handle_data(self, data):
        if self._ignorando:
            return
        if self._em_titulo:
            self.secao = (self.secao + " " + data.strip()).strip()[:255]
        self.atual.append(data)

    def _fechar(self):
        conteudo = re.sub(r"\n\s*\n+", "\n", "".join(self.atual)).strip()
        if conteudo:
            self.blocos.append(Bloco(None, self.secao, conteudo))
        self.atual = []


def _html(conteudo: str) -> list[Bloco]:
    extrator = _ExtratorHTML()
    extrator.feed(conteudo)
    extrator._fechar()
    return extrator.blocos
